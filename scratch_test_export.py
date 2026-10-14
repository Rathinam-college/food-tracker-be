import os
import django
from decimal import Decimal
from datetime import datetime

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'food_tracker_backend.settings')
django.setup()

from core.models import Item, StockBatch, VendorReturnItem
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

def test_export(item_id):
    try:
        item = Item.objects.get(id=item_id)
        print(f"Found item: {item.name}")
    except Item.DoesNotExist:
        print("Item not found")
        return

    batches = StockBatch.objects.filter(
        item_id=item_id,
        department__code='procurement'
    ).select_related(
        'item', 'department', 'procurement__purchase_request', 'procurement__vendor'
    ).prefetch_related(
        'allocations__mr_item__material_request'
    ).order_by('batch_date', 'created_at')

    print(f"Found {batches.count()} batches")

    vendor_returns = VendorReturnItem.objects.filter(
        item_id=item_id
    ).select_related('vendor_return', 'vendor_return__purchase_request')
    print(f"Found {vendor_returns.count()} returns")

    # Group returns by PR
    returns_by_pr = {}
    for vr_item in vendor_returns:
        pr_ref = vr_item.vendor_return.purchase_request.pr_reference
        if pr_ref not in returns_by_pr:
            returns_by_pr[pr_ref] = []
        returns_by_pr[pr_ref].append(vr_item)

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Batch Summary"
    
    # Just a subset of logic to check for crashes
    row = 4
    for b in batches:
        pr_ref = b.procurement.purchase_request.pr_reference if b.procurement and b.procurement.purchase_request else "N/A"
        b_issued = sum(a.quantity_allocated for a in b.allocations.all())
        b_returned = sum(r.quantity for r in returns_by_pr.get(pr_ref, [])) if pr_ref != "N/A" else Decimal('0')
        
        ws1.cell(row=row, column=1, value=b.batch_date)
        ws1.cell(row=row, column=2, value=pr_ref)
        row += 1

    ws2 = wb.create_sheet("Detailed Ledger")
    movements = []
    for b in batches:
        movements.append({
            'date': b.batch_date,
            'type': 'PURCHASE',
            'in': b.initial_quantity,
            'out': 0,
            'sort': b.created_at
        })
        for a in b.allocations.all():
            mr = a.mr_item.material_request
            movements.append({
                'date': mr.action_at.date() if mr.action_at else b.batch_date,
                'type': 'ISSUANCE',
                'in': 0,
                'out': a.quantity_allocated,
                'sort': mr.action_at if mr.action_at else b.created_at
            })
    
    movements.sort(key=lambda x: (x['date'], x['sort']))
    print("Sorted movements successfully")
    
    wb.save("test_output.xlsx")
    print("Saved workbook successfully")

if __name__ == "__main__":
    test_export(225)
