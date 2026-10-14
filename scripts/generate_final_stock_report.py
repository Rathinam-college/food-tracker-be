import os
import django
from decimal import Decimal
import openpyxl
from openpyxl.styles import Font, PatternFill

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "food_tracker_backend.settings")
django.setup()

from core.models import Item, PurchaseRequestItem, MaterialRequestItem, StockBatch

def generate_report():
    wb = openpyxl.Workbook()
    
    # Sheet 1: Purchases - MRs (Transaction Only)
    ws1 = wb.active
    ws1.title = "Transaction Only (Purch - MR)"
    
    # Sheet 2: With Adjustments
    ws2 = wb.create_sheet(title="With Adjustments")
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    
    headers1 = ["Item Name", "Purchased (In)", "Consumed (Out)", "Transaction Stock (=)"]
    headers2 = ["Item Name", "Purchased (In)", "Consumed (Out)", "Manual Adjustments (+)", "Total Stock (=)"]
    
    for i, header in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=i, value=header)
        cell.font = header_font
        cell.fill = header_fill
        
    for i, header in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=i, value=header)
        cell.font = header_font
        cell.fill = header_fill
        
    items = Item.objects.all().order_by('name')
    
    row = 2
    for item in items:
        # Calculate Purchases
        purchases = PurchaseRequestItem.objects.filter(
            item=item,
            quantity_received__gt=0
        ).exclude(purchase_request__status__in=['DRAFT', 'CANCELLED'])
        total_purchased = sum(p.quantity_received for p in purchases) or Decimal('0.00')
        
        # Calculate Consumed
        mrs = MaterialRequestItem.objects.filter(
            item=item,
            material_request__status__in=['ISSUED', 'ACCEPTED', 'PENDING_RECONCILE', 'COMPLETED']
        )
        total_consumed = Decimal('0.00')
        for mr in mrs:
            issued = mr.quantity_issued or Decimal('0.00')
            returned = mr.quantity_returned or Decimal('0.00')
            total_consumed += (issued - returned)
            
        # Calculate Manual Adjustments (StockBatch without procurement)
        adj_batches = StockBatch.objects.filter(item=item, procurement__isnull=True)
        total_adjusted = sum(b.initial_quantity for b in adj_batches) or Decimal('0.00')
        
        transaction_stock = total_purchased - total_consumed
        total_stock = transaction_stock + total_adjusted
        
        # Write Sheet 1
        ws1.cell(row=row, column=1, value=item.name)
        ws1.cell(row=row, column=2, value=float(total_purchased))
        ws1.cell(row=row, column=3, value=float(total_consumed))
        ws1.cell(row=row, column=4, value=float(transaction_stock)).font = Font(bold=True)
        
        # Write Sheet 2
        ws2.cell(row=row, column=1, value=item.name)
        ws2.cell(row=row, column=2, value=float(total_purchased))
        ws2.cell(row=row, column=3, value=float(total_consumed))
        ws2.cell(row=row, column=4, value=float(total_adjusted))
        ws2.cell(row=row, column=5, value=float(total_stock)).font = Font(bold=True)
        
        row += 1

    file_path = "d:/FoodDepartment/final_stock_report_with_adjustments.xlsx"
    wb.save(file_path)
    print(f"Report generated successfully at {file_path}")

if __name__ == '__main__':
    generate_report()
