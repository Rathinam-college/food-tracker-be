import os
import django
import openpyxl
from openpyxl.styles import Font, PatternFill

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "food_tracker_backend.settings")
try:
    django.setup()
except Exception:
    pass

from core.models import StockBatch

def generate_report():
    print("Finding all Ghost Batches (no procurement linked)...")
    ghost_batches = StockBatch.objects.filter(procurement__isnull=True).order_by('-batch_date', 'item__name')
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ghost Batches"
    
    headers = [
        "Batch ID", "Item ID", "Item Name", "Department", 
        "Initial Quantity", "Remaining Quantity", "Unit Price", "Batch Date", "Created At"
    ]
    
    ws.append(headers)
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        
    for b in ghost_batches:
        ws.append([
            b.id,
            b.item_id,
            b.item.name,
            b.department.name if b.department else "N/A",
            float(b.initial_quantity),
            float(b.remaining_quantity),
            float(b.unit_price),
            str(b.batch_date) if b.batch_date else "",
            str(b.created_at.date()) if b.created_at else ""
        ])
        
    for column_cells in ws.columns:
        length = max(len(str(cell.value) or "") for cell in column_cells)
        ws.column_dimensions[openpyxl.utils.get_column_letter(column_cells[0].column)].width = length + 2
        
    filepath = "d:/FoodDepartment/ghost_batches_report.xlsx"
    wb.save(filepath)
    print(f"Report saved to: {filepath}")
    print(f"Total ghost batches found: {ghost_batches.count()}")

if __name__ == '__main__':
    generate_report()
