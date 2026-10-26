import os
import django
from decimal import Decimal
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# Initialize Django environment
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "food_tracker_backend.settings")
try:
    django.setup()
except Exception:
    django.setup()

from core.models import MaterialRequest, StockMovement
from django.db.models import Sum

def generate_checker_excel():
    mrs = MaterialRequest.objects.filter(status__in=['ISSUED', 'ACCEPTED', 'PENDING_RECONCILE', 'COMPLETED']).order_by('-id')
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Anomaly Audit"
    
    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    ledger_header_font = Font(bold=True)
    ledger_header_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    
    current_row = 1
    mismatch_count = 0

    for mr in mrs:
        for mr_item in mr.items.all():
            item_name = mr_item.item.name
            mr_qty = mr_item.quantity_issued or Decimal('0.000')
            
            ledger_qty = StockMovement.objects.filter(
                reference_id=mr.mr_number,
                item=mr_item.item,
                type='STORE_TO_KITCHEN'
            ).aggregate(total=Sum('quantity'))['total'] or Decimal('0.000')

            if mr_qty != ledger_qty:
                mismatch_count += 1
                
                # --- MR ITEM DETAILS HEADER ---
                ws.cell(row=current_row, column=1, value="MR Number").font = header_font
                ws.cell(row=current_row, column=2, value="Item").font = header_font
                ws.cell(row=current_row, column=3, value="Requested Qty").font = header_font
                ws.cell(row=current_row, column=4, value="Issued Qty").font = header_font
                ws.cell(row=current_row, column=5, value="Returned Qty").font = header_font
                ws.cell(row=current_row, column=6, value="Deficit Qty").font = header_font
                for col in range(1, 7):
                    ws.cell(row=current_row, column=col).fill = header_fill
                current_row += 1
                
                # --- MR ITEM DETAILS VALUES ---
                ws.cell(row=current_row, column=1, value=mr.mr_number)
                ws.cell(row=current_row, column=2, value=item_name)
                ws.cell(row=current_row, column=3, value=float(mr_item.quantity))
                ws.cell(row=current_row, column=4, value=float(mr_item.quantity_issued))
                ws.cell(row=current_row, column=5, value=float(mr_item.quantity_returned))
                ws.cell(row=current_row, column=6, value=float(mr_item.deficit_quantity))
                current_row += 1
                
                # --- LEDGER HEADER ---
                ws.cell(row=current_row, column=1, value="STOCK MOVEMENT LEDGER").font = ledger_header_font
                ws.cell(row=current_row, column=2, value="ID").font = ledger_header_font
                ws.cell(row=current_row, column=3, value="Date").font = ledger_header_font
                ws.cell(row=current_row, column=4, value="User").font = ledger_header_font
                ws.cell(row=current_row, column=5, value="Qty").font = ledger_header_font
                ws.cell(row=current_row, column=6, value="Remarks").font = ledger_header_font
                for col in range(1, 7):
                    ws.cell(row=current_row, column=col).fill = ledger_header_fill
                current_row += 1
                
                # --- LEDGER ROWS ---
                movements = StockMovement.objects.filter(
                    reference_id=mr.mr_number,
                    item=mr_item.item,
                    type='STORE_TO_KITCHEN'
                ).order_by('id')
                
                for m in movements:
                    ws.cell(row=current_row, column=2, value=m.id)
                    ws.cell(row=current_row, column=3, value=m.movement_date.strftime('%Y-%m-%d %H:%M:%S'))
                    ws.cell(row=current_row, column=4, value=m.user.email if m.user else "System")
                    ws.cell(row=current_row, column=5, value=float(m.quantity))
                    ws.cell(row=current_row, column=6, value=m.remarks)
                    current_row += 1
                    
                # --- TOTALS ---
                ws.cell(row=current_row, column=4, value="Total Pushed:").font = Font(bold=True)
                ws.cell(row=current_row, column=5, value=float(ledger_qty)).font = Font(bold=True, color="FF0000")
                current_row += 2 # Add extra blank row for spacing
                
    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = adjusted_width

    output_path = 'd:/FoodDepartment/checker.xlsx'
    wb.save(output_path)
    print(f"Report generated successfully with {mismatch_count} anomalies at {output_path}")

if __name__ == '__main__':
    generate_checker_excel()
