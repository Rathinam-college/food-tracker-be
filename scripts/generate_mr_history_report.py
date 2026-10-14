import os
import django
from decimal import Decimal
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from collections import defaultdict

# Initialize Django environment
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "food_tracker_backend.settings")
try:
    django.setup()
except Exception:
    django.setup()

from core.models import MaterialRequestItem

def generate_mr_report():
    # Fetch all MR items ordered by item name and then MR target date
    # Only include MRs that have moved past Draft/Placed and are actually being processed
    mr_items = MaterialRequestItem.objects.select_related(
        'material_request', 'item'
    ).filter(
        material_request__status__in=['ISSUED', 'ACCEPTED', 'PENDING_RECONCILE', 'COMPLETED'],
        material_request__target_date__lte='2026-06-30'
    ).order_by('item__name', 'material_request__target_date')

    # Group by Item
    data_by_item = defaultdict(list)
    for mr_item in mr_items:
        data_by_item[mr_item.item.name].append(mr_item)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Material Request History"
    
    # Styles
    item_header_font = Font(bold=True, color="FFFFFF", size=12)
    item_header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    
    col_header_font = Font(bold=True)
    col_header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    
    total_font = Font(bold=True, color="FF0000")

    current_row = 1

    for item_name, items in data_by_item.items():
        # --- ITEM HEADER ROW ---
        ws.cell(row=current_row, column=1, value=f"ITEM: {item_name}").font = item_header_font
        for col in range(1, 9):
            ws.cell(row=current_row, column=col).fill = item_header_fill
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=8)
        current_row += 1
        
        # --- COLUMN HEADERS ---
        headers = [
            "MR Number", "Target Date", "Status", "Requested Qty", 
            "Issued Qty", "Returned Qty", "Consumed Qty", "Consumption Cost"
        ]
        for i, header in enumerate(headers, 1):
            cell = ws.cell(row=current_row, column=i, value=header)
            cell.font = col_header_font
            cell.fill = col_header_fill
        current_row += 1
        
        # --- MR ROWS ---
        total_requested = Decimal('0.00')
        total_issued = Decimal('0.00')
        total_returned = Decimal('0.00')
        total_consumed = Decimal('0.00')
        total_cost = Decimal('0.00')
        
        for mr_item in items:
            mr = mr_item.material_request
            
            req_qty = mr_item.quantity or Decimal('0.00')
            iss_qty = mr_item.quantity_issued or Decimal('0.00')
            ret_qty = mr_item.quantity_returned or Decimal('0.00')
            cons_qty = iss_qty - ret_qty
            cost = mr_item.consumption_total_cost or Decimal('0.00')
            
            ws.cell(row=current_row, column=1, value=mr.mr_number)
            ws.cell(row=current_row, column=2, value=mr.target_date.strftime('%Y-%m-%d') if mr.target_date else "")
            ws.cell(row=current_row, column=3, value=mr.status)
            ws.cell(row=current_row, column=4, value=float(req_qty))
            ws.cell(row=current_row, column=5, value=float(iss_qty))
            ws.cell(row=current_row, column=6, value=float(ret_qty))
            ws.cell(row=current_row, column=7, value=float(cons_qty))
            ws.cell(row=current_row, column=8, value=float(cost))
            
            total_requested += req_qty
            total_issued += iss_qty
            total_returned += ret_qty
            total_consumed += cons_qty
            total_cost += cost
            current_row += 1
            
        # --- TOTALS ROW ---
        ws.cell(row=current_row, column=3, value="TOTALS:").font = total_font
        ws.cell(row=current_row, column=4, value=float(total_requested)).font = total_font
        ws.cell(row=current_row, column=5, value=float(total_issued)).font = total_font
        ws.cell(row=current_row, column=6, value=float(total_returned)).font = total_font
        ws.cell(row=current_row, column=7, value=float(total_consumed)).font = total_font
        ws.cell(row=current_row, column=8, value=float(total_cost)).font = total_font
        
        current_row += 2 # Extra space between items
        
    widths = [20, 15, 18, 15, 15, 15, 15, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    output_path = 'd:/FoodDepartment/mr_history_report.xlsx'
    wb.save(output_path)
    print(f"Report generated successfully at {output_path}")

if __name__ == '__main__':
    generate_mr_report()
